#!/usr/bin/env python3
"""Build the complete CV from the supplied Word template with consistent styles.

Requires python-docx and the Codex documents renderer / LibreOffice.
Only the final PDF is a deliverable; the edited Word and page images are QA files.
"""
from pathlib import Path
from html import unescape
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.request
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT.parent / "tmp/cv-word-rebuilt"
OUTPUT = ROOT.parent / "output/pdf/Tomohide_R_Sato.pdf"
VERSION = "October 12, 2026"
LATIN = "Liberation Serif"
JAPANESE = "Noto Serif CJK JP"
BODY_SIZE = 9.5
BODY_LINE = 12.5


def text_lines(markup):
    text = re.sub(r"<br\s*/?>", "\n", markup, flags=re.I)
    text = unescape(re.sub(r"<[^>]*>", "", text))
    return [t.strip() for t in text.splitlines() if t.strip()]


def font(properties, name=LATIN, japanese=JAPANESE):
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for attr, value in [("ascii", name), ("hAnsi", name), ("cs", name), ("eastAsia", japanese)]:
        fonts.set(qn("w:" + attr), value)


def format_paragraph(paragraph, size, line, before=0, after=3):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = Pt(line)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.widow_control = True
    for run in paragraph._p.xpath(".//w:r"):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(0, properties)
        font(properties)
        for name in ["sz", "szCs"]:
            element = properties.find(qn("w:" + name))
            if element is None:
                element = OxmlElement("w:" + name)
                properties.append(element)
            element.set(qn("w:val"), str(round(size * 2)))


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    doc = Document(ROOT / "content/cv-template.docx")
    data = json.loads((ROOT / "content/editorial.json").read_text())
    en = next(s["entries"] for s in data["locales"]["en"]["cv"] if s["title"] == "Academic conferences")
    ja = next(s["entries"] for s in data["locales"]["ja"]["cv"] if s["title"] == "学会発表")
    for p in list(doc.paragraphs):
        if "Publication-related links:" in "".join(p._p.xpath(".//w:t/text()")):
            p._p.getparent().remove(p._p)
    anchor = next(p for p in doc.paragraphs if p.text.startswith("2026 Mar") and "World Congress" in p.text)
    for english, japanese in zip(en[:3], ja[:3], strict=True):
        p = Paragraph(OxmlElement("w:p"), anchor._parent)
        anchor._p.addprevious(p._p)
        p.style = "CV Entry"
        meeting, title = text_lines(english["html"])[:2]
        original_title = text_lines(japanese["html"])[1]
        p.add_run(english["date"] + "  ").bold = True
        p.add_run(meeting)
        p.add_run("\n    " + title + "\n    Japanese title: " + original_title)
    for node in doc.element.xpath("//w:t"):
        if node.text:
            node.text = node.text.replace("June 3, 2026", VERSION).replace(" (scheduled/listed on source page)", "")
    for style in doc.styles:
        if style.type == 1:
            font(style.element.get_or_add_rPr())
    for border in doc.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ["Normal", "CV Entry"]:
        style = doc.styles[name]
        style.font.size = Pt(BODY_SIZE)
        style.paragraph_format.line_spacing = Pt(BODY_LINE)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    for index, p in enumerate(doc.paragraphs):
        if index == 0:
            p.style = "Title"
            for border in p._p.xpath(".//w:pBdr"):
                border.getparent().remove(border)
            format_paragraph(p, 18, 22, after=3)
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif index == 1:
            format_paragraph(p, 14, 18, after=4)
        elif index == 2:
            format_paragraph(p, 9, 12, after=6)
        elif p.style.name == "Heading 1":
            format_paragraph(p, 12.5, 16, before=8, after=4)
            p.paragraph_format.keep_with_next = True
        elif p.style.name == "Heading 2":
            format_paragraph(p, 10.5, 14, before=6, after=3)
            p.paragraph_format.keep_with_next = True
        else:
            format_paragraph(p, BODY_SIZE, BODY_LINE)
            p.paragraph_format.keep_together = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for node in p._p.xpath(".//w:t"):
                        if node.text:
                            node.text = node.text.replace("Doctoral Fellow A", "Doctoral\u00a0Fellow\u00a0A")
                    format_paragraph(p, 8.5, 11, after=1)
    for section in doc.sections:
        for p in section.footer.paragraphs:
            for node in p._p.xpath(".//w:t"):
                node.text = node.text.replace("June 3, 2026", VERSION)
            format_paragraph(p, 8, 10, after=0)
    doc.core_properties.title = "Curriculum Vitae - Tomohide R. Sato"
    doc.core_properties.subject = "Curriculum vitae"
    working_docx = WORK / "Tomohide_R_Sato.docx"
    doc.save(working_docx)
    fonts_dir = WORK / "fonts"
    fonts_dir.mkdir(exist_ok=True)
    japanese_font = fonts_dir / "NotoSerifCJKjp-Regular.otf"
    if not japanese_font.exists():
        url = "https://raw.githubusercontent.com/notofonts/noto-cjk/main/Serif/OTF/Japanese/NotoSerifCJKjp-Regular.otf"
        with urllib.request.urlopen(url, timeout=30) as response:
            data = response.read()
        if not data.startswith(b"OTTO"):
            raise ValueError("Invalid Japanese font download")
        japanese_font.write_bytes(data)
    runtime = Path(os.environ["CODEX_PRIMARY_RUNTIME_ROOT"])
    config = fonts_dir / "fonts.conf"
    config.write_text(f'''<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
<dir>{runtime}/dependencies/native/libreoffice-headless/libreoffice/share/fonts/truetype</dir>
<dir>{runtime}/dependencies/native/poppler/poppler/fonts</dir>
<dir>{fonts_dir}</dir><cachedir>{WORK}/font-cache</cachedir>
</fontconfig>''')
    env = dict(os.environ, FONTCONFIG_FILE=str(config))
    render_dir = WORK / "render"
    subprocess.run([sys.executable, "/root/.codex/skills/builtins/documents/render_docx.py", str(working_docx),
                    "--output_dir", str(render_dir), "--emit_pdf"], check=True, env=env)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(render_dir / "Tomohide_R_Sato.pdf", OUTPUT)
    print(f"Rendered complete Word-based CV: {OUTPUT}")


if __name__ == "__main__":
    main()
